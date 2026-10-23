from pathlib import Path
import json
from docx import Document

ROOT = Path(".").resolve()
DATA = json.loads((ROOT / "work" / "corpus_doc" / "chars1000_part2.json").read_text(encoding="utf-8"))
expected = [r["char"] for r in DATA["chars"]]

doc = Document(ROOT / "docs" / "字伴-现代汉语常用2000字表-第2册.docx")
print("paragraphs", len(doc.paragraphs), "tables", len(doc.tables))
char_tables = [t for t in doc.tables if len(t.rows) == 10 and len(t.columns) == 10]
print("10x10 tables", len(char_tables))
actual = []
for t in char_tables:
    for row in t.rows:
        for cell in row.cells:
            parts = cell.text.strip().split()
            if len(parts) != 2:
                raise SystemExit(f"bad cell {cell.text!r}")
            actual.append(parts[1])

print("actual", len(actual), "unique", len(set(actual)))
print("order matches", actual == expected)
if actual != expected:
    for i, (a, b) in enumerate(zip(actual, expected)):
        if a != b:
            print("first mismatch", i, a, b)
            break

all_text = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(
    c.text for t in doc.tables for r in t.rows for c in r.cells
)
for needle in ("第 1001–2000", "96.7%", "99.2%", "必要性评分", "不是另一套独立字表"):
    print("contains", needle, needle in all_text)

print("first 20", "".join(actual[:20]))
print("last 20", "".join(actual[-20:]))
