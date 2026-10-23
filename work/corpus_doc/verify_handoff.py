from pathlib import Path
from docx import Document
import re

ROOT = Path(".").resolve()
docx_path = ROOT / "docs" / "字伴-语料批处理与自动审核-交接手册.docx"
md_path = ROOT / "docs" / "字伴-语料批处理与自动审核-交接手册.md"

doc = Document(docx_path)
headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
text = "\n".join(p.text for p in doc.paragraphs)
text += "\n" + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)

print("docx paragraphs", len(doc.paragraphs), "tables", len(doc.tables), "headings", len(headings))
print("first headings:", headings[:8])
print("last headings:", headings[-8:])

required = [
    "0. 30 秒快速开始",
    "1. 你接手的是什么",
    "2. 不可谈判的边界",
    "3. 输入：生词清单和来源包",
    "4. 数据格式：四个文件",
    "5. 批量流水线：从词到 release",
    "6. 硬 lint 规则",
    "7. 评审 rubric 与判定",
    "8. 校准：没有人工，不等于没有标准",
    "9. 运行时集成与冒烟",
    "10. 交付合同",
    "11. 常见坑",
    "12. 单人两周执行表",
    "附录 A · 生成提示词",
    "附录 B · Judge A 提示词",
    "附录 C · Judge B 提示词",
    "附录 D · Judge C / 三审提示词",
    "附录 E · 老人视角对抗提示词",
    "附录 F · 词表与覆盖数据",
    "附录 G · 交接检查清单",
]
missing = [h for h in required if h not in text]
print("missing headings", missing)

phrases = [
    "默认丢弃",
    "只处理我给的词",
    "hard_reject",
    "source 以 auto- 开头",
    "0 重复",
    "99.16%",
    "没有人工审校",
]
for p in phrases:
    print("phrase", p, p in text)

md = md_path.read_text(encoding="utf-8")
print("md chars", len(md), "lines", md.count("\n") + 1)
print("code fences", md.count("```"), "even", md.count("```") % 2 == 0)
print("md headings", len(re.findall(r"^#{1,3} ", md, re.M)))
