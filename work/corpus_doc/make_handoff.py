# -*- coding: utf-8 -*-
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from handoff_content import BLOCKS

ROOT = Path(".").resolve()
OUT_DOCX = ROOT / "docs" / "字伴-语料批处理与自动审核-交接手册.docx"
OUT_MD = ROOT / "docs" / "字伴-语料批处理与自动审核-交接手册.md"

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
BLUE = RGBColor(0x2E, 0x5C, 0x8A)
GRAY = RGBColor(0x66, 0x66, 0x66)
LIGHT = RGBColor(0x8C, 0x8C, 0x8C)
RED = RGBColor(0xA6, 0x2A, 0x2A)

BODY_FONT = "宋体"
HEAD_FONT = "微软雅黑"
MONO_FONT = "Consolas"


def set_run_font(run, name=BODY_FONT, size=10.5, bold=False, color=None, italic=False, eastasia=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), eastasia or name)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)


def add_para(text="", size=10.5, bold=False, color=None, align=None, space_after=6, space_before=0, line=1.25):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing = line
    if text:
        run = p.add_run(text)
        set_run_font(run, size=size, bold=bold, color=color)
    return p


def add_heading(text, level=1):
    h = doc.add_heading(level=level)
    h.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    set_run_font(
        run,
        name=HEAD_FONT,
        size=17 if level == 1 else 12.5,
        bold=True,
        color=NAVY if level == 1 else BLUE,
    )
    return h


def add_bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.2
    run = p.add_run(text)
    set_run_font(run, size=10)
    return p


def add_numbered(number, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.first_line_indent = Cm(-0.5)
    p.paragraph_format.line_spacing = 1.2
    r = p.add_run(f"{number}. ")
    set_run_font(r, size=10, bold=True, color=BLUE)
    r2 = p.add_run(text)
    set_run_font(r2, size=10)
    return p


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(table, top=40, start=60, bottom=40, end=60):
    tbl_pr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tbl_pr.append(mar)


def set_table_borders(table, color="D9D9D9", size=4):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tbl_pr.append(borders)


def keep_row_together(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)
    set_run_font(run, size=9, color=GRAY)


def add_table(headers, rows, widths=None):
    ncols = len(headers)
    if not widths:
        widths = [16.8 / ncols] * ncols
    table = doc.add_table(rows=len(rows) + 1, cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, color="D9D9D9", size=4)
    set_cell_margins(table, top=45, start=65, bottom=45, end=65)

    for j, header in enumerate(headers):
        cell = table.cell(0, j)
        cell.width = Cm(widths[j])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(cell, "EAF0F7")
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_run_font(r, size=9, bold=True, color=NAVY)

    for i, row in enumerate(rows, start=1):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.width = Cm(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            r = p.add_run(str(value))
            set_run_font(r, size=8.6, bold=(j == 0 and ncols <= 2), color=RGBColor(0x22, 0x22, 0x22))

    for row in table.rows:
        keep_row_together(row)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_code(text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, color="E5E5E5", size=4)
    set_cell_margins(table, top=65, start=90, bottom=65, end=90)
    cell = table.cell(0, 0)
    cell.width = Cm(16.8)
    keep_row_together(table.rows[0])
    shade(cell, "F7F7F7")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    for i, line in enumerate(text.splitlines()):
        if i:
            p.add_run().add_break()
        run = p.add_run(line)
        set_run_font(run, name=MONO_FONT, size=8.1, color=RGBColor(0x2B, 0x2B, 0x2B), eastasia="宋体")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(1.9)
section.bottom_margin = Cm(1.7)
section.left_margin = Cm(2.0)
section.right_margin = Cm(2.0)
section.header_distance = Cm(0.8)
section.footer_distance = Cm(0.8)

normal = doc.styles["Normal"]
normal.font.name = BODY_FONT
normal.font.size = Pt(10.5)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
normal.paragraph_format.line_spacing = 1.25

for style_name, size in (("Heading 1", 17), ("Heading 2", 12.5), ("Heading 3", 11)):
    style = doc.styles[style_name]
    style.font.name = HEAD_FONT
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = NAVY if style_name == "Heading 1" else BLUE
    style._element.rPr.rFonts.set(qn("w:eastAsia"), HEAD_FONT)

header_p = section.header.paragraphs[0]
header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hr = header_p.add_run("字伴 · 语料批处理与自动审核 · 交接手册")
set_run_font(hr, size=8, color=LIGHT)
add_page_number(section.footer.paragraphs[0])

# cover
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(62)
r = p.add_run("字伴 · 内容基础设施")
set_run_font(r, name=HEAD_FONT, size=10.5, bold=True, color=BLUE)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(12)
p.paragraph_format.space_after = Pt(2)
r = p.add_run("语料批处理与自动审核")
set_run_font(r, name=HEAD_FONT, size=28, bold=True, color=NAVY)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(20)
r = p.add_run("交接手册 · 给下一个 AI")
set_run_font(r, name=HEAD_FONT, size=13, color=GRAY)

meta_rows = [
    ("版本", "v1.0 · 2026-10-23"),
    ("适用", "下一个接手语料批处理的 AI / 维护者"),
    ("输入", "生词清单、来源包、1–2000 字表"),
    ("输出", "corpus/release/*.json + 评审记录 + 批次报告"),
    ("硬约束", "无人工审校；默认丢弃；机器审校必须标注"),
]
table = doc.add_table(rows=len(meta_rows), cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
set_table_borders(table, color="E5E5E5", size=4)
set_cell_margins(table, top=65, start=90, bottom=65, end=90)
for i, (key, value) in enumerate(meta_rows):
    table.columns[0].width = Cm(3.6)
    table.columns[1].width = Cm(13.2)
    c0, c1 = table.cell(i, 0), table.cell(i, 1)
    c0.width = Cm(3.6)
    c1.width = Cm(13.2)
    for c in (c0, c1):
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    shade(c0, "F2F5F9")
    p0 = c0.paragraphs[0]
    p0.paragraph_format.space_after = Pt(0)
    r0 = p0.add_run(key)
    set_run_font(r0, size=9.5, color=GRAY)
    p1 = c1.paragraphs[0]
    p1.paragraph_format.space_after = Pt(0)
    r1 = p1.add_run(value)
    set_run_font(r1, size=9.5, bold=True, color=RGBColor(0x22, 0x22, 0x22))

add_para("", space_after=12)
add_para(
    "这份手册不承诺“机器审校等于人工审校”。它承诺的是：在单人、无人工、无广泛试验对象的条件下，"
    "仍然把明显错误挡在发布之外，并把不确定性如实告诉用户。",
    size=9.5,
    color=GRAY,
    align=WD_ALIGN_PARAGRAPH.LEFT,
)

pagebreak_count = 0
for idx, block in enumerate(BLOCKS):
    is_last = idx == len(BLOCKS) - 1
    kind = block[0]
    if kind == "h1":
        add_heading(block[1], 1)
    elif kind == "h2":
        add_heading(block[1], 2)
    elif kind == "h3":
        add_heading(block[1], 3)
    elif kind == "p":
        add_para(block[1], space_after=6)
    elif kind == "note":
        add_para(block[1], size=9.5, color=GRAY, space_before=2, space_after=6)
    elif kind == "bullets":
        for item in block[1]:
            add_bullet(item)
        if not is_last:
            add_para("", space_after=2)
    elif kind == "numbers":
        for i, item in enumerate(block[1], start=1):
            add_numbered(i, item)
        if not is_last:
            add_para("", space_after=2)
    elif kind == "table":
        add_table(block[1], block[2], block[3] if len(block) > 3 else None)
    elif kind == "code":
        add_code(block[1])
    elif kind == "pagebreak":
        doc.add_page_break()
        pagebreak_count += 1

doc.core_properties.title = "字伴语料批处理与自动审核交接手册"
doc.core_properties.subject = "给下一个 AI 的语料批量生成、自动评审、安装与验收手册"
doc.core_properties.author = "字伴项目"
doc.core_properties.language = "zh-CN"
doc.core_properties.comments = "无人工审校模式；硬规则 + 多模型评审 + 默认丢弃。"

OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT_DOCX)
print("saved", OUT_DOCX)
print("paragraphs", len(doc.paragraphs), "tables", len(doc.tables), "pagebreaks", pagebreak_count)


# ---------------- markdown ----------------
def md_escape(text):
    return str(text).replace("|", "\\|")


lines = []
for block in BLOCKS:
    kind = block[0]
    if kind == "h1":
        lines += [f"# {block[1]}", ""]
    elif kind == "h2":
        lines += [f"## {block[1]}", ""]
    elif kind == "h3":
        lines += [f"### {block[1]}", ""]
    elif kind == "p":
        lines += [block[1], ""]
    elif kind == "note":
        lines += [f"> {block[1]}", ""]
    elif kind == "bullets":
        lines += [f"- {item}" for item in block[1]]
        lines.append("")
    elif kind == "numbers":
        lines += [f"{i}. {item}" for i, item in enumerate(block[1], start=1)]
        lines.append("")
    elif kind == "table":
        headers, rows = block[1], block[2]
        lines.append("| " + " | ".join(md_escape(h) for h in headers) + " |")
        lines.append("|" + "|".join(["---"] * len(headers)) + "|")
        for row in rows:
            lines.append("| " + " | ".join(md_escape(v) for v in row) + " |")
        lines.append("")
    elif kind == "code":
        lines += ["```text", block[1], "```", ""]
    elif kind == "pagebreak":
        lines += ["---", ""]

md = "\n".join(lines).rstrip() + "\n"
OUT_MD.write_text(md, encoding="utf-8")
print("saved", OUT_MD, "chars", len(md))
