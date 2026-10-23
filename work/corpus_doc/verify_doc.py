from pathlib import Path
import json
from docx import Document

ROOT = Path(".").resolve()
DATA = json.loads((ROOT / "work" / "corpus_doc" / "chars1000.json").read_text(encoding="utf-8"))
expected = [r["char"] for r in DATA["chars"]]

doc = Document(ROOT / "docs" / "字伴-现代汉语常用1000字表.docx")
print("paragraphs", len(doc.paragraphs), "tables", len(doc.tables))
print("sections", len(doc.sections))

# Char tables are the 10 consecutive 10x10 tables after cover and replacement.
char_tables = [t for t in doc.tables if len(t.rows) == 10 and len(t.columns) == 10]
print("10x10 tables", len(char_tables))
actual = []
for t in char_tables:
    for row in t.rows:
        for cell in row.cells:
            text = cell.text.strip()
            # text is like "0001 的"
            parts = text.split()
            if len(parts) != 2:
                raise SystemExit(f"bad cell {text!r}")
            actual.append(parts[1])

print("actual chars", len(actual), "unique", len(set(actual)))
print("order matches", actual == expected)
if actual != expected:
    for i, (a, b) in enumerate(zip(actual, expected)):
        if a != b:
            print("first mismatch", i, a, b)
            break
print("first 20", "".join(actual[:20]))
print("last 20", "".join(actual[-20:]))

# Headings and key text
headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
print("headings", len(headings))
for h in headings[:20]:
    print("H:", h)

all_text = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
checks = [
    "现代汉语常用 1000 字表",
    "89.0%",
    "默认丢弃",
    "机器审校",
    "不能证明",
    "近代汉语",
]
for c in checks:
    print("contains", c, c in all_text)
