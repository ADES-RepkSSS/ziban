from pathlib import Path
import json

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(".").resolve()
DATA = json.loads((ROOT / "work" / "corpus_doc" / "chars1000_part2.json").read_text(encoding="utf-8"))
CHARS = DATA["chars"]
META = DATA["meta"]
PROJECT = META["project_sample"]
OUT = ROOT / "docs" / "字伴-现代汉语常用2000字表-第2册.docx"

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
BLUE = RGBColor(0x2E, 0x5C, 0x8A)
GRAY = RGBColor(0x66, 0x66, 0x66)
LIGHT = RGBColor(0x8C, 0x8C, 0x8C)
RED = RGBColor(0xA6, 0x2A, 0x2A)

BODY_FONT = "宋体"
HEAD_FONT = "微软雅黑"
MONO_FONT = "Consolas"


def set_run_font(run, name=BODY_FONT, size=10.5, bold=False, color=None, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)


def add_para(text="", size=10.5, bold=False, color=None, align=None, space_after=6, space_before=0, line=1.25, name=BODY_FONT):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing = line
    if text:
        run = p.add_run(text)
        set_run_font(run, name=name, size=size, bold=bold, color=color)
    return p


def add_heading(text, level=1):
    h = doc.add_heading(level=level)
    h.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    set_run_font(run, name=HEAD_FONT, size=16 if level == 1 else 12.5, bold=True, color=NAVY if level == 1 else BLUE)
    return h


def add_bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.2
    if bold_prefix:
        r = p.add_run(bold_prefix)
        set_run_font(r, size=10.5, bold=True)
    r = p.add_run(text)
    set_run_font(r, size=10.5)
    return p


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(table, top=30, start=40, bottom=30, end=40):
    tbl_pr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tbl_pr.append(mar)


def set_table_borders(table, color="D9D9D9", size=4):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tbl_pr.append(borders)


def keep_row_together(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def add_hyperlink(paragraph, url, text):
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    r_pr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    r_pr.append(u)
    new_run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)
    set_run_font(run, size=9, color=GRAY)


def add_code_block(text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, color="E5E5E5", size=4)
    set_cell_margins(table, top=70, start=100, bottom=70, end=100)
    cell = table.cell(0, 0)
    cell.width = Cm(16.8)
    shade(cell, "F7F7F7")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.1
    for i, line in enumerate(text.splitlines()):
        if i:
            p.add_run().add_break()
        run = p.add_run(line)
        set_run_font(run, name=MONO_FONT, size=8.5, color=RGBColor(0x33, 0x33, 0x33))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(2.0)
section.bottom_margin = Cm(1.8)
section.left_margin = Cm(2.1)
section.right_margin = Cm(2.1)
section.header_distance = Cm(0.9)
section.footer_distance = Cm(0.9)

normal = doc.styles["Normal"]
normal.font.name = BODY_FONT
normal.font.size = Pt(10.5)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
normal.paragraph_format.line_spacing = 1.25

for style_name, size in (("Heading 1", 16), ("Heading 2", 12.5), ("Heading 3", 11)):
    style = doc.styles[style_name]
    style.font.name = HEAD_FONT
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = NAVY if style_name == "Heading 1" else BLUE
    style._element.rPr.rFonts.set(qn("w:eastAsia"), HEAD_FONT)

header_p = section.header.paragraphs[0]
header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hr = header_p.add_run("字伴 · 现代汉语常用 2000 字表 · 第 2 册")
set_run_font(hr, size=8, color=LIGHT)
add_page_number(section.footer.paragraphs[0])

# cover
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(72)
r = p.add_run("字伴 · 内容基础设施")
set_run_font(r, name=HEAD_FONT, size=10.5, bold=True, color=BLUE)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(10)
p.paragraph_format.space_after = Pt(2)
r = p.add_run("现代汉语常用 2000 字表")
set_run_font(r, name=HEAD_FONT, size=27, bold=True, color=NAVY)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(18)
r = p.add_run("第 2 册 · 第 1001–2000 字 · 频率序 · 限定于《现代汉语常用字表》2500 字")
set_run_font(r, name=HEAD_FONT, size=11, color=GRAY)

meta_rows = [
    ("本册", "1000 个唯一汉字（序号 1001–2000）"),
    ("累计", "2000 字；第 1 册 + 第 2 册，0 重复"),
    ("本册新增覆盖", f"{META['part2_marginal_coverage']*100:.1f}% 的 Jun Da 语料 token"),
    ("累计覆盖", f"{META['cumulative_coverage']*100:.1f}%"),
    ("官方 2500 剩余", f"{META['official_remaining']} 个未收入"),
    ("项目卡片覆盖", f"{PROJECT['coverage']*100:.1f}%（50 张卡释义 + 例句，共 {PROJECT['cjk_chars']} 字）"),
]
table = doc.add_table(rows=len(meta_rows), cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
set_table_borders(table, color="E5E5E5", size=4)
set_cell_margins(table, top=70, start=100, bottom=70, end=100)
for i, (key, value) in enumerate(meta_rows):
    table.columns[0].width = Cm(4.2)
    table.columns[1].width = Cm(12.6)
    c0, c1 = table.cell(i, 0), table.cell(i, 1)
    c0.width = Cm(4.2)
    c1.width = Cm(12.6)
    for c in (c0, c1):
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    shade(c0, "F2F5F9")
    p0 = c0.paragraphs[0]
    p0.paragraph_format.space_after = Pt(0)
    r0 = p0.add_run(key)
    set_run_font(r0, size=9.5, color=GRAY)
    p1 = c1.paragraphs[0]
    p1.paragraph_format.space_after = Pt(0)
    r1 = p1.add_run(value)
    set_run_font(r1, size=9.5, bold=True, color=RGBColor(0x22, 0x22, 0x22))

add_para("", space_after=14)
add_para(
    "这是第 1 册的续册，不是另一套独立字表。第 1 册是 1–1000，第 2 册是 1001–2000；两册合起来才是 2000 字。",
    size=9.5,
    color=GRAY,
    space_after=4,
)
add_para(
    "第 2 册的字更靠后、更书面、更偏专名。它适合做“释义用字上限”的扩展，不代表这些字都要教，也不代表按这个顺序教。",
    size=9.5,
    color=GRAY,
)

# summary
add_heading("摘要", 1)
add_bullet("同一方法：Jun Da 现代汉语字频表排序，只保留《现代汉语常用字表》2500 字，取第 1001–2000 个唯一汉字；第 1 册 1000 字不变。", "口径不变：")
add_bullet(f"第 2 册新增 {META['part2_marginal_coverage']*100:.1f}% token 覆盖；两册合计 {META['cumulative_coverage']*100:.1f}%。收益明显递减，这是正常的。", "覆盖变化：")
add_bullet(f"50 张卡的释义/例句用字覆盖从 95.7% 升到 {PROJECT['coverage']*100:.1f}%；只剩 {len(PROJECT['unknown_unique'])} 个表外字（{('、'.join(PROJECT['unknown_unique']))}）。", "项目检查：")
add_bullet("第 2 册的字更适合“允许出现”，不适合“强制使用”。释义仍应优先用第 1 册；第 2 册只在需要时兜底。", "使用原则：")

# 1
add_heading("1. 续接口径", 1)
add_para(
    "第 1 册的取字过程是：频率排序 → 限定官方 2500 → 取前 1000。第 2 册完全沿用，只是从第 1001 个继续取到第 2000 个。"
    "第 1 册里被替换掉的 7 个字（尔、伊、谓、诺、伦、俄、洛）不会在第 2 册复活；第 2 册也不重复前 1000 字。",
    space_after=6,
)
add_para(
    f"两册合计 2000 个唯一汉字，0 重复；官方 2500 字里还剩 {META['official_remaining']} 个未收入。"
    "剩下的 500 字频率更低，包含更多专名、书面语、行业字；继续扩到 2500 的边际收益很小，不建议为了“收满”而收。",
    space_after=6,
)

# 2 coverage
add_heading("2. 覆盖变化：第 2 册买到了什么", 1)
add_para(
    "第 1 册覆盖 89.0%，第 2 册把累计覆盖推到 96.7%。新增的 7.6 个百分点，换来的是“更少卡在生僻字上”，"
    "但也意味着第 2 册里大量字在日常生活中并不常见。",
    space_after=4,
)
coverage_rows = [
    ("1000", "89.0%", "第 1 册结束"),
    ("1100", "90.5%", ""),
    ("1200", "91.7%", ""),
    ("1300", "92.7%", ""),
    ("1400", "93.6%", ""),
    ("1500", "94.3%", "官方 2500 的 60%"),
    ("1600", "94.9%", ""),
    ("1700", "95.5%", ""),
    ("1800", "95.9%", ""),
    ("1900", "96.3%", ""),
    ("2000", "96.7%", "第 2 册结束"),
]
t = doc.add_table(rows=len(coverage_rows) + 1, cols=3)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
t.autofit = False
set_table_borders(t, color="D9D9D9", size=4)
set_cell_margins(t, top=40, start=60, bottom=40, end=60)
for j, header in enumerate(("累计字数", "累计 token 覆盖", "说明")):
    cell = t.cell(0, j)
    cell.width = [Cm(3.4), Cm(4.8), Cm(8.6)][j]
    shade(cell, "EAF0F7")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j < 2 else WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(header)
    set_run_font(r, size=9.5, bold=True, color=NAVY)
for i, row in enumerate(coverage_rows, start=1):
    for j, value in enumerate(row):
        cell = t.cell(i, j)
        cell.width = [Cm(3.4), Cm(4.8), Cm(8.6)][j]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j < 2 else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(value)
        set_run_font(r, size=9, bold=(j == 1), color=RGBColor(0x22, 0x22, 0x22))

add_para(
    "读法：从 1000 到 2000，累计覆盖只增加 7.6 个百分点；从 1500 到 2000 只增加 2.4 个百分点。"
    "所以第 2 册的定位是“减少生僻字卡顿”，不是“多认 1000 个字”。",
    size=9.5,
    color=GRAY,
    space_before=4,
    space_after=6,
)

# 3 table
doc.add_page_break()
add_heading("3. 第 1001–2000 字", 1)
add_para(
    "排列规则：按频率从高到低；编号是 2000 字表内的序号（1001–2000），不是原始频率排名；每字只出现一次。"
    "带拼音、频次和原始频率排名的版本见 chars1000_part2.csv。",
    size=9.5,
    color=GRAY,
    space_after=8,
)
cum = dict(META["cumulative_by_100"])
for block in range(10):
    start = 1000 + block * 100
    end = start + 100
    if block in (2, 4, 6, 8):
        doc.add_page_break()
    add_heading(f"第 {start + 1}–{end} 字（累计覆盖 {cum[end]:.1f}%）", 2)
    table = doc.add_table(rows=10, cols=10)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, color="D9D9D9", size=4)
    set_cell_margins(table, top=28, start=20, bottom=28, end=20)
    col_width = Cm(1.66)
    for row in table.rows:
        row.height = Cm(0.78)
        keep_row_together(row)
        for cell in row.cells:
            cell.width = col_width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
    for idx in range(block * 100, (block + 1) * 100):
        item = CHARS[idx]
        r_i, c_i = divmod(idx - block * 100, 10)
        cell = table.cell(r_i, c_i)
        p = cell.paragraphs[0]
        rank_run = p.add_run(f"{1001 + idx:04d}")
        set_run_font(rank_run, size=6.5, color=LIGHT)
        char_run = p.add_run(" " + item["char"])
        set_run_font(char_run, name="宋体", size=13.5, color=RGBColor(0x11, 0x11, 0x11))
        if block % 2 == 1:
            shade(cell, "FAFAFA")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

# 4 usage
doc.add_page_break()
add_heading("4. 第 2 册怎么用", 1)
add_heading("4.1 释义用字：先第 1 册，后第 2 册", 2)
add_para(
    "第 2 册的正确用法不是“把 2000 字都塞进 prompt”，而是给释义一个更宽的上限：",
    space_after=3,
)
for text in (
    "默认仍用第 1 册 1000 字写释义；第 2 册的字只在无法替代时出现。",
    "释义中 95% 以上的汉字落在 1–2000 字内；低于 90% 直接拒绝。",
    "第 2 册的字连续出现不超过 2 个；出现时要能解释为什么非用不可。",
    "2001–2500 的字只允许出现在目标词本身、地名/机构名/药名等专名里，不用于释义。",
):
    add_bullet(text)

add_heading("4.2 目标词仍然看屏幕证据，不看字表", 2)
add_para(
    "第 2 册里有很多字，能组成词，但不值得收。比如“篇、凡、诸、倘、蔽、颠”这类字，老人屏幕上很少遇到，"
    "或者遇到也不影响办事。不要因为字进了 2000 字表，就给它配一个词条。",
    space_after=3,
)
for text in (
    "这个词会不会出现在老人的屏幕、牌子、单据、App 上？",
    "看不懂它，会不会让老人办不成事、按错、被骗？",
    "有没有更简单的词已经覆盖了同一个场景？",
):
    add_bullet(text)
add_para(
    "三个问题有一个答不上来，就不收。第 2 册是“允许出现的字”，不是“必须生产的词条”。",
    bold=True,
    color=NAVY,
    space_before=3,
    space_after=6,
)

add_heading("4.3 自动审核新增的三条规则", 2)
add_para(
    "第 1 册文档第 5 节的流水线不变：硬规则、3 稿、双评审、三审、默认丢弃。第 2 册只加三条针对“低频字”的规则。",
    space_after=4,
)
rules = [
    ("层级检查", "释义用字按 1–1000 / 1001–2000 / 2001–2500 分三层；第 3 层只允许词内字/专名。"),
    ("生僻词惩罚", "如果存在更简单的同义词，评审应拒绝；不能用“书面词换口语词”来显得准确。"),
    ("必要性评分", "在五项 rubric 外增加第六项“值不值得收”（1–5）；低于 4 分直接丢弃，不看其他分数。"),
]
t = doc.add_table(rows=len(rules) + 1, cols=2)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
t.autofit = False
set_table_borders(t, color="D9D9D9", size=4)
set_cell_margins(t, top=50, start=70, bottom=50, end=70)
for j, header in enumerate(("规则", "判定")):
    cell = t.cell(0, j)
    cell.width = Cm(3.2) if j == 0 else Cm(13.6)
    shade(cell, "EAF0F7")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(header)
    set_run_font(r, size=9.5, bold=True, color=NAVY)
for i, (name, desc) in enumerate(rules, start=1):
    for j, value in enumerate((name, desc)):
        cell = t.cell(i, j)
        cell.width = Cm(3.2) if j == 0 else Cm(13.6)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(value)
        set_run_font(r, size=9, bold=(j == 0), color=RGBColor(0x22, 0x22, 0x22))

add_para(
    "锚点集也要加 20 条：10 条“字很靠后但词很该收”（如“账、频、缴”），10 条“字很靠后但词不值得收”（如“倘、蔽、颠”）。"
    "如果评审器分不清这两类，它就会把低频词当高价值内容。",
    size=9.5,
    color=GRAY,
    space_before=4,
    space_after=6,
)

doc.add_page_break()
add_heading("5. 单人执行：第 2 册不是产量目标", 1)
add_para(
    "第 1 册的 300 条语料目标不需要因为第 2 册翻倍。第 2 册只是把“释义可用的字”放宽，"
    "不是让你再生成 1000 条词条。真正要扩语料时，仍然按屏幕证据选词，按第 1 册的自动审核流程走。",
    space_after=3,
)
for text in (
    "先扩到 400–500 条，不直接冲 1000 条；每批 150 条，自动审核后可能只剩 90–110 条。",
    "第 2 册的字进入释义时，让 lint 打 flag；flag 率超过 20% 就说明你在用低频字硬写。",
    "每批加 10 条“低频但该收 / 低频但不该收”的锚点，防止评审器偏好书面语。",
    "不要对外说“2000 字覆盖率 96.7%”等于“老人能读懂 96.7% 的屏幕”；它只是语料 token 覆盖率。",
):
    add_bullet(text)

# 6 sources
doc.add_page_break()
add_heading("6. 来源、文件与复现", 1)
add_para("方法与第 1 册一致，来源也相同：", space_after=3)
for label, url in (
    ("Jun Da, Modern Chinese Character Frequency List", "https://lingua.mtsu.edu/chinese-computing/statistics/char/list.php?Which=MO"),
    ("《现代汉语常用字表》2500 常用字（可获取副本）", "http://basic.nmg.smartedu.cn/player/data2/2020/11/12/21/2/5EB279DB5BF27AEEE053F1C8A8C0E792/a76c25774cf0ec35d2b385bc4085c754.docx"),
):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Cm(0.4)
    r = p.add_run("· " + label + "：")
    set_run_font(r, size=9.5)
    add_hyperlink(p, url, url)

add_para(
    "配套文件：chars1000_part2.csv / json / txt 是第 2 册；chars2000.csv / json / txt 是两册合并版。",
    size=9,
    color=GRAY,
    space_before=4,
    space_after=6,
)

# appendix A
add_heading("附录 A · 第 2 册前 30 与末 30", 1)
add_para("前 30：" + "".join(r["char"] for r in CHARS[:30]), size=10, space_after=3)
add_para("末 30：" + "".join(r["char"] for r in CHARS[-30:]), size=10, space_after=6)

# appendix B
add_heading("附录 B · 项目卡片用字变化", 1)
add_para(
    f"第 1 册覆盖 95.66%，表外字 21 个；第 2 册加入后覆盖 {PROJECT['coverage']*100:.2f}%，"
    f"表外字只剩 {len(PROJECT['unknown_unique'])} 个，共出现 {PROJECT['unknown_count']} 次。",
    size=9.5,
    space_after=3,
)
add_para("剩余表外字：" + "、".join(PROJECT["unknown_unique"]), size=9, color=GRAY, space_after=6)
add_para(
    "这说明：2000 字已经足够覆盖现有卡片语言；再往上扩到 2500，对现有内容的帮助很小。"
    "后续更应该把精力放在“选对词、写清楚、自动审核”上，而不是继续加字。",
    size=9.5,
    color=NAVY,
    bold=True,
)

add_para("", space_after=8)
add_para(
    "第 2 册的价值不是“多 1000 个字”，而是让你在写释义时少一次“这个字不在表里”的中断；"
    "它不应该变成新的内容产量指标。",
    size=10,
    color=GRAY,
    align=WD_ALIGN_PARAGRAPH.CENTER,
)

doc.core_properties.title = "现代汉语常用 2000 字表 · 第 2 册"
doc.core_properties.subject = "第 1001–2000 字；字伴屏幕词库的释义用字扩展"
doc.core_properties.author = "字伴项目"
doc.core_properties.language = "zh-CN"
doc.core_properties.comments = "与第 1 册同一口径：Jun Da 字频表 + 《现代汉语常用字表》2500 字过滤。"

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print("saved", OUT)
print("paragraphs", len(doc.paragraphs), "tables", len(doc.tables))
